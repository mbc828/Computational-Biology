# Public source snapshot: replace path placeholders before running.
from pathlib import Path
import json
import re
import numpy as np
import openpyxl

ROOT = Path(r"<PROJECT_ROOT>\outputs\manuscript_v21_github_release\review_v88")
SRC = ROOT.parent.parent / "elisa_redraw_20261026" / "source"
DOCK = ROOT / "source" / "docking_local_run_20260923T2336"

specs = [
    ("IL-1β", SRC / "IL-1b.xlsx", "D40:I40", "D41:I41", (7.716, 20.272, 0.5275)),
    ("IL-6", SRC / "IL-6.xlsx", "C42:H42", "C43:H43", (0.5599, 39.067, 3.1522)),
    ("TNF-α", SRC / "TNF-a.xlsx", "D42:I42", "D43:I43", (32.194, 275.45, 3.9883)),
]

def values(ws, ref):
    return [c.value for row in ws[ref] for c in row]

def fit(x, y):
    x = np.asarray(x, dtype=float)
    y = np.asarray(y, dtype=float)
    coef = np.polyfit(x, y, 2)
    pred = np.polyval(coef, x)
    r2 = 1.0 - np.sum((y-pred)**2) / np.sum((y-y.mean())**2)
    return coef, r2, pred

report = {"elisa": {}, "docking": {}}
lines = []
for name, path, xrange, yrange, workbook_coef in specs:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    x = values(ws, xrange)
    y = values(ws, yrange)
    coef, r2, pred = fit(x, y)
    unit = ws[chr(ord(xrange[0])) + str(int(xrange[1:3]) + 1)].value if False else None
    entry = {
        "file": str(path),
        "sheet": ws.title,
        "od_cells": xrange,
        "concentration_cells": yrange,
        "od_means": x,
        "concentrations": y,
        "fit_coefficients_full_precision": coef.tolist(),
        "r2_full_precision": float(r2),
        "predicted_concentrations": pred.tolist(),
        "residuals": (np.asarray(y)-pred).tolist(),
        "workbook_sample_formula_coefficients": workbook_coef,
        "rounded_coefficients_match_workbook": np.allclose(np.round(coef, 4), workbook_coef, atol=1e-12),
        "r2_rounded_4": round(float(r2), 4),
    }
    report["elisa"][name] = entry
    lines += [f"{name}: C={coef[0]:.4f}*OD^2+{coef[1]:.4f}*OD+{coef[2]:.4f}; R2={r2:.8f}; R2(4dp)={r2:.4f}", f"  OD={x}", f"  concentrations={y}", f"  residuals={(np.asarray(y)-pred).tolist()}", f"  coefficients match workbook after 4dp rounding: {entry['rounded_coefficients_match_workbook']}"]

logs = sorted(DOCK.glob("*/vina.log"))
for log in logs:
    text = log.read_text(encoding="utf-8")
    def grab(pattern):
        m = re.search(pattern, text)
        return m.group(1).strip() if m else None
    modes = re.findall(r"^\s*(\d+)\s+(-?\d+(?:\.\d+)?)\s+", text, re.M)
    entry = {
        "version": grab(r"AutoDock Vina v([^\r\n]+)"),
        "receptor": grab(r"Rigid receptor:\s*([^\r\n]+)"),
        "ligand": grab(r"Ligand:\s*([^\r\n]+)"),
        "grid_center": grab(r"Grid center:\s*([^\r\n]+)"),
        "grid_size": grab(r"Grid size\s*:\s*([^\r\n]+)"),
        "grid_space": grab(r"Grid space\s*:\s*([^\r\n]+)"),
        "exhaustiveness": grab(r"Exhaustiveness:\s*([^\r\n]+)"),
        "cpu": grab(r"CPU:\s*([^\r\n]+)"),
        "seed": grab(r"Performing docking \(random seed:\s*(\d+)\)"),
        "num_modes_observed": len(modes),
        "best_affinity_kcal_mol": float(modes[0][1]) if modes else None,
    }
    report["docking"][log.parent.name] = entry
    lines.append(f"DOCK {log.parent.name}: {entry}")

(ROOT / "verified_parameters.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
(ROOT / "verified_parameters.txt").write_text("\n".join(lines), encoding="utf-8")
print("\n".join(lines[:18]))
print(f"Verified {len(logs)} Vina logs")
