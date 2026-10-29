# Public source snapshot: replace path placeholders before running.
from pathlib import Path
from copy import copy
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import DataBarRule
from openpyxl.utils import get_column_letter

ROOT = Path(r'<PROJECT_ROOT>/outputs/manuscript_v21_github_release')
SRC = ROOT / 'review_v86/source/体外抗菌.xlsx'
OUT = ROOT / 'SHXD_PJI_抗菌数据核对与统计_v86.xlsx'
source = load_workbook(SRC, data_only=False)
ss = source['Sheet1']

wb = Workbook()
ws = wb.active
ws.title = '原始数据'
raw = wb.create_sheet('全组汇总')
bai = wb.create_sheet('黄芩素核对')
notes = wb.create_sheet('说明与来源')

navy = '17365D'
blue = 'D9EAF7'
light = 'F2F2F2'
white = 'FFFFFF'
red = 'C00000'
thin_gray = Side(style='thin', color='D9E1F2')

for sh in wb.worksheets:
    sh.sheet_view.showGridLines = False
    sh.freeze_panes = 'B2'

# Raw transcription: retain exactly what the source records, including blanks.
headers = [ss.cell(1, c).value for c in range(1, 18)]
ws.append(headers)
for r in range(2, 5):
    ws.append([ss.cell(r, c).value for c in range(1, 18)])
ws.append(['来源说明'] + ['体外抗菌.xlsx / Sheet1；A2:A4 为复板编号，空白表示原表未记录。'] + [''] * 15)
ws.append(['原表换算规则'] + ['CFU/mL = 菌落数 × 10^6（原表注明：稀释 5 次、取 0.1 mL）'] + [''] * 15)
ws.append(['浓度单位'] + ['原表未注明；64/32/16 等仅为原表浓度标识，本文档不推定 μM 或 μg/mL。'] + [''] * 15)
ws.append(['原表备注'] + ['对照（旧）标注“有问题”；复板均值按实际非空复板计算。'] + [''] * 15)
for cell in ws[1]:
    cell.fill = PatternFill('solid', fgColor=navy)
    cell.font = Font(color=white, bold=True)
    cell.alignment = Alignment(wrap_text=True, vertical='center')
ws.row_dimensions[1].height = 42
ws.column_dimensions['A'].width = 18
for c in range(2, 18):
    ws.column_dimensions[get_column_letter(c)].width = 16
for r in range(6, 10):
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=17)
    ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(r, 1).font = Font(bold=True, color=navy)
    ws.row_dimensions[r].height = 32

# All-group formula summary. Empty source cells are not converted to zero.
raw_headers = ['组别', '原表浓度标识', '复板1', '复板2', '复板3', '有效复板数', '菌落均值', '菌落SD', '均值CFU/mL', '记录完整性']
raw.append(raw_headers)
for idx, c in enumerate(range(2, 18), start=2):
    header = str(ss.cell(1, c).value)
    if '（' in header:
        name, level = header.rsplit('（', 1)
        level = level.rstrip('）')
    else:
        name, level = header, ''
    letter = get_column_letter(c)
    row = idx
    raw.cell(row, 1, name)
    raw.cell(row, 2, level)
    for j, rr in enumerate(range(2, 5), start=3):
        raw.cell(row, j, f'=IF(\'原始数据\'!{letter}{rr}="","",\'原始数据\'!{letter}{rr})')
    raw.cell(row, 6, f'=COUNT(C{row}:E{row})')
    raw.cell(row, 7, f'=IF(F{row}>0,AVERAGE(C{row}:E{row}),"")')
    raw.cell(row, 8, f'=IF(F{row}>1,STDEV.S(C{row}:E{row}),"")')
    raw.cell(row, 9, f'=IF(F{row}>0,G{row}*1000000,"")')
    raw.cell(row, 10, f'=IF(F{row}=3,"3个复板",IF(F{row}=0,"无复板记录","仅"&F{row}&"个复板"))')
raw['A19'] = '说明'
raw['B19'] = '均值与 SD 均按非空复板计算；仅 1 个复板时不计算 SD。CFU/mL 使用原表 ×10^6 规则。'
raw.merge_cells('B19:J19')

# Dedicated, traceable baicalein check using the new control only.
bai_headers = ['组别', '原表浓度标识', '复板1', '复板2', '复板3', 'n', '菌落均值', 'SD', 'CFU/mL均值', '相对新对照存活率(%)']
bai.append(bai_headers)
selected = [('对照（新）', 'Q'), ('黄芩素（16）', 'M'), ('黄芩素（32）', 'L'), ('黄芩素（64）', 'K')]
for row, (label, col) in enumerate(selected, start=2):
    bai.cell(row, 1, label)
    bai.cell(row, 2, label.split('（')[-1].rstrip('）') if '（' in label else '')
    for j, rr in enumerate(range(2, 5), start=3):
        bai.cell(row, j, f'=IF(\'原始数据\'!{col}{rr}="","",\'原始数据\'!{col}{rr})')
    bai.cell(row, 6, f'=COUNT(C{row}:E{row})')
    bai.cell(row, 7, f'=AVERAGE(C{row}:E{row})')
    bai.cell(row, 8, f'=STDEV.S(C{row}:E{row})')
    bai.cell(row, 9, f'=G{row}*1000000')
    bai.cell(row, 10, f'=G{row}/$G$2*100')
bai['A7'] = '核对结论'
bai['B7'] = '黄芩素三档均有 3 个非空复板；原表 64 档三复板均为 0。浓度单位、复板独立性及与图 6 单浓度实验的对应关系须由实验记录确认，未将两组数据混合。'
bai.merge_cells('B7:J7')
bai['A8'] = '统计口径'
bai['B8'] = '本表只作描述性复算；未将复板自动认定为独立生物学重复，未新增显著性检验。'
bai.merge_cells('B8:J8')

chart = BarChart()
chart.type = 'col'
chart.style = 10
chart.title = '黄芩素原表浓度标识与平板计数'
chart.y_axis.title = 'CFU/mL（按原表规则换算）'
chart.x_axis.title = '组别（浓度单位未注明）'
chart.height = 8.2
chart.width = 15.5
data = Reference(bai, min_col=9, min_row=1, max_row=5)
cats = Reference(bai, min_col=1, min_row=2, max_row=5)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.legend = None
bai.add_chart(chart, 'A11')

notes_rows = [
    ('项目', 'SHXD–PJI 抗菌数据核对与统计 v86'),
    ('原始文件', '体外抗菌.xlsx（2026-10-29 附件复制件）'),
    ('原始工作表', 'Sheet1；Sheet2、Sheet3 为空'),
    ('数据来源', '用户提供的实验记录；未使用图片估算菌落数。'),
    ('换算规则', '原表注明：菌落数 × 10^6 = CFU/mL（稀释 5 次、取 0.1 mL）。'),
    ('可用记录', '黄芩素 16/32/64 三档与新对照均有 3 个非空复板，可做均值、SD、CFU/mL 和相对新对照存活率的描述性复算。'),
    ('未确认字段', '原表未注明浓度单位；复板是否独立生物学重复、实验日期、批次及与既有单浓度图的对应关系未记录。'),
    ('排除规则', '旧对照由原表标注“有问题”；不用于新对照归一化。空白复板不按 0 处理；单复板组不计算 SD。'),
    ('正文处理', 'v86 正文图 6 保留既有可追溯的平板计数和活/死染色结果；本新表暂不与旧数据合并，待单位与批次归属确认后再补入。'),
]
for row in notes_rows:
    notes.append(row)
notes.column_dimensions['A'].width = 20
notes.column_dimensions['B'].width = 110
for r in range(1, notes.max_row + 1):
    notes.cell(r, 1).font = Font(bold=True, color=navy)
    notes.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
    notes.row_dimensions[r].height = 35

for sh in (raw, bai):
    for cell in sh[1]:
        cell.fill = PatternFill('solid', fgColor=navy)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sh.row_dimensions[1].height = 38
    widths = [22, 16, 12, 12, 12, 10, 14, 14, 18, 24]
    for i, width in enumerate(widths, start=1):
        sh.column_dimensions[get_column_letter(i)].width = width
    for row in sh.iter_rows(min_row=2, max_row=5 if sh is bai else 17):
        for cell in row:
            cell.border = Border(bottom=thin_gray)
            cell.alignment = Alignment(vertical='center')
    for r in range(2, 6 if sh is bai else 18):
        sh.cell(r, 7).number_format = '0.00'
        sh.cell(r, 8).number_format = '0.00'
        sh.cell(r, 9).number_format = '0.00E+00'
        if sh is bai:
            sh.cell(r, 10).number_format = '0.00'
    sh.auto_filter.ref = f'A1:J{5 if sh is bai else 17}'

bai.conditional_formatting.add('I2:I5', DataBarRule(start_type='num', start_value=0, end_type='max', color='5B9BD5', showValue=True))
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = 'auto'
wb.save(OUT)
print(OUT)
