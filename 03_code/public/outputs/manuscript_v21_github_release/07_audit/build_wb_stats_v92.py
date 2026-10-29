# Public source snapshot: replace path placeholders before running.
from pathlib import Path
import json

import numpy as np
from openpyxl import load_workbook
from scipy import stats

ROOT = Path(r"<PROJECT_ROOT>")
SOURCE = ROOT / "outputs" / "manuscript_v21_github_release" / "06_source_material" / "wetlab_wb_raw_20261030" / "10.14灰度值(1).xlsx"
WORK = ROOT / "outputs" / "wet_figures_20261014" / "work"
OLD = WORK / "stats.json"
OUT = WORK / "stats_v92.json"

wb = load_workbook(SOURCE, data_only=True, read_only=True)
ws = wb["Sheet1"]
columns = {1: (3, 4, 5), 2: (7, 8, 9), 3: (11, 12, 13)}
rows = {"pIkBa": 2, "IkBa": 3, "pp65": 5, "p65": 6, "nuc": 8, "cyto": 9}
raw = {}
for rep, cols in columns.items():
    for gi, group in enumerate(("Control", "PJI+SHXD", "PJI")):
        for protein, row in rows.items():
            raw[(rep, group, protein)] = float(ws.cell(row, cols[gi]).value)
wb.close()

groups = ("Control", "PJI+SHXD", "PJI")
metrics = {
    "p-IkBa/IkBa": lambda rep, group: raw[(rep, group, "pIkBa")] / raw[(rep, group, "IkBa")],
    "p-p65/p65": lambda rep, group: raw[(rep, group, "pp65")] / raw[(rep, group, "p65")],
    "Nuclear/Cytoplasmic p65": lambda rep, group: raw[(rep, group, "nuc")] / raw[(rep, group, "cyto")],
    "Nuclear p65 fraction": lambda rep, group: raw[(rep, group, "nuc")] / (raw[(rep, group, "nuc")] + raw[(rep, group, "cyto")]),
    "Cytoplasmic p65 fraction": lambda rep, group: raw[(rep, group, "cyto")] / (raw[(rep, group, "nuc")] + raw[(rep, group, "cyto")]),
}

result = json.loads(OLD.read_text(encoding="utf-8"))
result["WB"] = {}
print("WB v92 statistics: two-sided equal-variance Student t tests; raw P × 3 Bonferroni")
for metric, fn in metrics.items():
    block = {"groups": {}}
    arrays = {}
    for group in groups:
        values = [float(fn(rep, group)) for rep in (1, 2, 3)]
        arr = np.asarray(values, dtype=float)
        arrays[group] = arr
        block["groups"][group] = {
            "values": values,
            "mean": float(np.mean(arr)),
            "sd": float(np.std(arr, ddof=1)),
            "n": 3,
        }
    f_stat, anova_p = stats.f_oneway(*(arrays[group] for group in groups))
    block["anova"] = {"F": float(f_stat), "p": float(anova_p)}
    block["pairwise"] = {}
    for a, b in (("Control", "PJI+SHXD"), ("Control", "PJI"), ("PJI+SHXD", "PJI")):
        test = stats.ttest_ind(arrays[a], arrays[b], equal_var=True)
        p_adj = min(1.0, float(test.pvalue) * 3)
        block["pairwise"][f"{a} vs {b}"] = {
            "t": float(test.statistic),
            "df": 4,
            "p_raw": float(test.pvalue),
            "p": p_adj,
            "adjustment": "Bonferroni raw P × 3",
        }
        print(f"{metric} | {a} vs {b}: t={float(test.statistic):.9f}; P_adj={p_adj:.12g}")
    result["WB"][metric] = block

result["WB_statistics"] = {
    "source": SOURCE.name,
    "ratio_values": "Calculated within each biological replicate from the archived grey-value workbook.",
    "summary": "Mean and sample SD; n=3 per group.",
    "omnibus": "One-way ANOVA across three groups.",
    "pairwise": "Two-sided equal-variance Student t tests with Bonferroni adjustment (raw P × 3).",
    "note": "Replaces legacy pairwise entries that contained unadjusted Welch P values; group means and SDs are unchanged.",
}
OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"WROTE={OUT}")
