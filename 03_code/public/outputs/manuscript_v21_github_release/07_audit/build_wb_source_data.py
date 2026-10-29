# Public source snapshot: replace the path placeholder before running.
# -*- coding: utf-8 -*-
"""Build the submission-oriented WB Source Data file-correspondence workbook."""
from __future__ import annotations

import hashlib
import json
import zlib
from pathlib import Path

import rarfile
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(r"<PROJECT_ROOT>\outputs\manuscript_v21_github_release")
RAW = ROOT / "06_source_material" / "wetlab_wb_raw_20261030"
AUDIT = ROOT / "07_audit" / "wb_source_mapping"
OUT = ROOT / "04_data" / "wb_audit" / "WB_Source_Data文件对应清单_v1_20261030.xlsx"
OUT.parent.mkdir(parents=True, exist_ok=True)
SOURCE = RAW / "10.14灰度值(1).xlsx"
GROUPS = ["Control", "HKSA+Baicalein", "HKSA"]
ROWS = {"p-IκBα": 2, "IκBα": 3, "p-p65": 5, "p65": 6,
        "Nuclear p65": 8, "Cytoplasmic p65": 9, "β-actin": 10}
COLUMNS = {1: (3, 4, 5), 2: (7, 8, 9), 3: (11, 12, 13)}
STRIPS = {
    "p-IκBα": "022.pptx / image1.png",
    "IκBα": "022.pptx / image4.png",
    "p-p65": "044.pptx / image1.png",
    "p65": "044.pptx / image4.png",
    "Nuclear p65": "022.pptx / image7.png",
    "Cytoplasmic p65": "022.pptx / image10.png",
    "β-actin": "022.pptx / image13.png",
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def crc32(path: Path) -> str:
    h = 0
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h = zlib.crc32(chunk, h)
    return f"{h & 0xffffffff:08x}"


def target_for(name: str) -> str:
    lower = name.lower()
    if lower.startswith("p-ikba"):
        return "p-IκBα"
    if lower.startswith("ikba"):
        return "IκBα"
    if lower.startswith("p-p65") or lower.startswith("pp65"):
        return "p-p65"
    if "p65核" in lower:
        return "Nuclear p65"
    if "p65质" in lower:
        return "Cytoplasmic p65"
    if lower.startswith("p65"):
        return "p65"
    if lower.startswith("actin"):
        return "β-actin"
    raise ValueError(f"Unrecognized target: {name}")


analysis = json.loads((AUDIT / "image_provenance_analysis.json").read_text(encoding="utf-8"))
inventory = analysis["inventory"]
assert len(inventory) == 27

# The archive is the original delivery package tying 044.pptx to six numbered TIFFs.
rar_rows = []
primary_from_rar = set()
with rarfile.RarFile(RAW / "10144(1).rar") as archive:
    for entry in archive.infolist():
        if entry.isdir() or not entry.filename.lower().endswith(".tif"):
            continue
        basename = Path(entry.filename).name
        rep = int(basename.split("(", 1)[1][0])
        rel = f"rep{rep}/{basename}"
        local = RAW / rel
        local_crc = crc32(local)
        archive_crc = f"{entry.CRC:08x}"
        match = entry.file_size == local.stat().st_size and entry.CRC == int(local_crc, 16)
        if not match:
            raise ValueError(f"RAR/local mismatch: {rel}")
        primary_from_rar.add(rel)
        rar_rows.append({
            "entry": entry.filename,
            "local": rel,
            "bytes": entry.file_size,
            "archive_crc32": archive_crc,
            "local_crc32": local_crc,
            "match": "一致",
            "archive_time": "-".join(f"{x:02d}" for x in entry.date_time[:3]) + " " + ":".join(f"{x:02d}" for x in entry.date_time[3:6]),
        })
assert len(rar_rows) == 6

ppt_evidence = []
for image_name in ("image1.png", "image4.png"):
    candidates = [x for x in analysis["template_matches"] if x["ppt_image"] == image_name]
    candidates.sort(key=lambda x: -x.get("score", -1))
    best = candidates[0]
    same_target = [x for x in candidates if Path(x["candidate"]).name.lower().startswith(("p-p65" if image_name == "image1.png" else "p65"))]
    second = same_target[1] if len(same_target) > 1 else candidates[1]
    ppt_evidence.append({
        "entry": f"044.pptx / ppt/media/{image_name}",
        "local": best["candidate"],
        "score": best["score"],
        "next_score": second["score"],
        "next_candidate": second["candidate"],
        "scale": best["scale"],
    })

# Non-duplicated target files are primary by one-to-one target/replicate mapping;
# numbered p65/p-p65 files are primary by the RAR package evidence.
records = []
for item in inventory:
    rel = item["path"]
    name = Path(rel).name
    rep = int(Path(rel).parent.name.replace("rep", ""))
    target = target_for(name)
    is_p65_pair = target in {"p65", "p-p65"}
    primary = rel in primary_from_rar if is_p65_pair else True
    row = ROWS[target]
    columns = COLUMNS[rep]
    cells = [f"{get_column_letter(c)}{row}" for c in columns]
    duplicate_group = f"rep{rep}_{target}" if is_p65_pair else ""
    if primary:
        role = "主用原膜"
        if is_p65_pair:
            status = "主用文件确认；ROI记录待补"
            evidence = "10144(1).rar 与 044.pptx 同源打包；大小及 CRC32 一致"
            if rep == 1:
                match = next(x for x in ppt_evidence if x["local"] == rel)
                evidence += f"；PPT 裁带 NCC={match['score']:.6f}"
        else:
            status = "唯一主文件；ROI记录待补"
            evidence = "该重复／靶标仅有这一份全膜文件"
        figure_use = STRIPS[target] if rep == 1 else "图 5 仅展示 rep1 裁带"
        gray_mapping = "主用文件对应"
    else:
        role = "备选版本"
        status = "备选保留；不作灰度来源指认"
        evidence = "未纳入 10144(1).rar 的 044 来源包"
        figure_use = "未用于图 5"
        gray_mapping = "同名靶标单元格，仅用于识别候选"
    rar = next((x for x in rar_rows if x["local"] == rel), None)
    records.append({
        **item,
        "rep": rep,
        "target": target,
        "role": role,
        "duplicate_group": duplicate_group,
        "gray_row": f"Sheet1!{row}",
        "gray_cells": " / ".join(cells),
        "group_order": " / ".join(GROUPS),
        "gray_mapping": gray_mapping,
        "figure_use": figure_use,
        "evidence": evidence,
        "status": status,
        "remaining": "原始 ImageJ ROI、背景扣除与曝光选择记录" if primary else "无需删除；归档保留以维持来源完整",
        "rar": rar,
    })

assert sum(x["role"] == "主用原膜" for x in records) == 21
assert sum(x["role"] == "备选版本" for x in records) == 6

# Read all 63 source grey values directly, without recalculating densitometry.
source_wb = load_workbook(SOURCE, data_only=True, read_only=True)
source_ws = source_wb["Sheet1"]
gray_records = []
primary_lookup = {(x["rep"], x["target"]): x for x in records if x["role"] == "主用原膜"}
for target, row in ROWS.items():
    for rep in (1, 2, 3):
        for gi, group in enumerate(GROUPS):
            cell = f"{get_column_letter(COLUMNS[rep][gi])}{row}"
            value = source_ws[cell].value
            if not isinstance(value, (int, float)):
                raise ValueError(f"Missing grey value: {cell}")
            primary = primary_lookup[(rep, target)]
            gray_records.append({
                "target": target,
                "rep": rep,
                "group": group,
                "cell": f"Sheet1!{cell}",
                "value": value,
                "path": primary["path"],
                "file_status": primary["status"],
                "trace": "灰度值逐格转录；具体 ROI 与背景扣除参数未归档",
            })
source_wb.close()
assert len(gray_records) == 63

# Workbook styling helpers.
wb = Workbook()
wb.remove(wb.active)
NAVY = "244F70"
BLUE = "EAF1F6"
GREY = "F3F5F7"
GREEN = "E7F2EA"
AMBER = "FFF4D6"
WHITE = "FFFFFF"
TEXT = "17232D"
MEDIUM = Side(style="medium", color=NAVY)
HAIR = Side(style="hair", color="D9DEE2")


def setup(ws, widths: dict[str, float], freeze="B4"):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = freeze
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_view.zoomScale = 90


def title(ws, text: str, end_col: int, subtitle: str | None = None):
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=end_col)
    c = ws.cell(1, 2, text)
    c.font = Font(name="Arial", size=16, bold=True, color=NAVY)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    if subtitle:
        ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=end_col)
        c = ws.cell(2, 2, subtitle)
        c.font = Font(name="Arial", size=10, italic=True, color="52616D")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[2].height = 28


def headers(ws, row: int, values: list[str], start=2):
    for offset, value in enumerate(values):
        c = ws.cell(row, start + offset, value)
        c.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=MEDIUM)
    ws.row_dimensions[row].height = 32


def body(ws, min_row: int, max_row: int, min_col: int, max_col: int, row_height=30):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for c in row:
            c.font = Font(name="Arial", size=10, color=TEXT)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            c.border = Border(bottom=HAIR)
        ws.row_dimensions[row[0].row].height = row_height


def add_table_sheet(name: str, sheet_title: str, subtitle: str, header_values: list[str], rows: list[list], widths: dict[str, float]):
    ws = wb.create_sheet(name)
    setup(ws, widths)
    end_col = 1 + len(header_values)
    title(ws, sheet_title, end_col, subtitle)
    headers(ws, 3, header_values)
    for ri, values in enumerate(rows, 4):
        for ci, value in enumerate(values, 2):
            ws.cell(ri, ci, value)
    body(ws, 4, 3 + len(rows), 2, end_col)
    ws.auto_filter.ref = f"B3:{get_column_letter(end_col)}{3 + len(rows)}"
    return ws


# Summary sheet with formula-derived counts.
ws = wb.create_sheet("摘要")
setup(ws, {"A": 3, "B": 30, "C": 24, "D": 24, "E": 24, "F": 24, "G": 24, "H": 24}, freeze="B3")
title(ws, "WB Source Data 文件对应清单", 8, "版本 v1｜2026-10-30｜覆盖 27 个 TIFF、63 个原始灰度单元格及 044 来源包证据")
summary_metrics = [
    ("TIFF 文件总数", "=COUNTA('文件对应清单'!B4:B30)", "27 个文件均保留"),
    ("主用原膜", "=COUNTIF('文件对应清单'!G4:G30,\"主用原膜\")", "7 个靶标 × 3 次重复"),
    ("p65／p-p65 主用文件", "=COUNTIFS('文件对应清单'!G4:G30,\"主用原膜\",'文件对应清单'!D4:D30,\"p65\")+COUNTIFS('文件对应清单'!G4:G30,\"主用原膜\",'文件对应清单'!D4:D30,\"p-p65\")", "三批带编号 TIFF 均由 044 来源包确认"),
    ("备选版本", "=COUNTIF('文件对应清单'!G4:G30,\"备选版本\")", "不带编号文件及 rep3 小图，不删除"),
    ("044 包内原膜", "=COUNTIF('文件对应清单'!X4:X30,\"一致\")", "文件大小与 CRC32 逐项一致"),
]
for r, (label, formula, note) in enumerate(summary_metrics, 4):
    ws.cell(r, 2, label).font = Font(name="Arial", bold=True, color=NAVY)
    ws.cell(r, 3, formula).font = Font(name="Arial", size=12, bold=True, color=TEXT)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
    ws.cell(r, 4, note).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 26

conclusions = [
    ("主用 p65／p-p65 原膜", "rep1／rep2／rep3 分别使用带 (1)／(2)／(3) 编号的 TIFF。10144(1).rar 同时包含 044.pptx 和这 6 个 TIFF；包内条目与本地文件的大小、CRC32 全部一致。"),
    ("rep1 图件来源", "044.pptx 的 p-p65 裁带对应 rep1/P-P65 (1).tif，NCC=0.996505；p65 裁带对应 rep1/P65 (1).tif，NCC=0.998103。"),
    ("备选文件处置", "rep1／rep2 不带编号的 p65、p-p65，以及 rep3 的 p65.tif、pp65.tif 小图均作为备选版本保留，未纳入 044 主用来源包，不与现稿灰度行建立直接来源指认。"),
    ("灰度追溯边界", "63 个灰度值已逐格映射至主用文件。当前资料未包含原始 ImageJ ROI、背景扣除及曝光选择记录，因此本清单确认文件级对应，不声称复现了原始定量 ROI。"),
    ("数据与统计", "本工作簿不重新圈选条带、不改变任何灰度值、比率、P 值或内参；内参继续按 β-actin 归档。"),
]
start = 11
for r, (label, text) in enumerate(conclusions, start):
    ws.cell(r, 2, label).font = Font(name="Arial", bold=True, color=NAVY)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    ws.cell(r, 3, text).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 50
ws.cell(18, 2, "工作表导航").font = Font(name="Arial", bold=True, color=NAVY)
navigation = [
    ("文件对应清单", "27 个 TIFF 的角色、灰度行、PPT／RAR 证据和 SHA-256"),
    ("灰度单元格映射", "63 个灰度值逐格对应至主用原膜"),
    ("PPT与RAR证据", "6 个 CRC32 校验及 2 个 PPT 图像匹配记录"),
    ("来源与口径", "来源文件哈希、状态定义及使用边界"),
]
for r, (label, text) in enumerate(navigation, 19):
    ws.cell(r, 2, label).font = Font(name="Arial", bold=True, color=TEXT)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    ws.cell(r, 3, text)
    ws.row_dimensions[r].height = 26

file_headers = ["文件ID", "重复", "靶标", "文件名", "相对路径", "文件角色", "同名候选组",
                "灰度表行", "灰度单元格（Control / HKSA+Baicalein / HKSA）", "灰度映射口径",
                "图 5 rep1 裁带", "对应证据", "确认状态", "待补／归档说明", "字节数", "像素尺寸",
                "模式", "DPI", "Photoshop 标签", "SHA-256", "RAR CRC32", "本地 CRC32", "CRC 结果"]
file_rows = []
for index, item in enumerate(records, 1):
    rar = item["rar"]
    file_rows.append([
        f"WB-TIFF-{index:03d}", f"rep{item['rep']}", item["target"], Path(item["path"]).name, item["path"], item["role"],
        item["duplicate_group"], item["gray_row"], item["gray_cells"], item["gray_mapping"], item["figure_use"],
        item["evidence"], item["status"], item["remaining"], item["bytes"], f"{item['width']}×{item['height']}",
        item["mode"], item["dpi_x"], "有" if item["has_photoshop_tag_37724"] else "无", item["sha256"],
        rar["archive_crc32"] if rar else None, rar["local_crc32"] if rar else None, rar["match"] if rar else None,
    ])
ws = add_table_sheet("文件对应清单", "27 个 TIFF 文件对应清单",
                     "主用原膜用于现稿文件级映射；备选版本完整保留。灰度单元格列的组序均为 Control / HKSA+Baicalein / HKSA。",
                     file_headers, file_rows,
                     {"A": 3, "B": 15, "C": 10, "D": 19, "E": 23, "F": 32, "G": 13, "H": 19, "I": 13,
                      "J": 31, "K": 24, "L": 25, "M": 48, "N": 30, "O": 38, "P": 15, "Q": 14, "R": 10,
                      "S": 10, "T": 16, "U": 68, "V": 15, "W": 15, "X": 12})
ws.conditional_formatting.add("N4:N30", FormulaRule(formula=['ISNUMBER(SEARCH("确认",N4))'], fill=PatternFill("solid", fgColor=GREEN)))
ws.conditional_formatting.add("N4:N30", FormulaRule(formula=['ISNUMBER(SEARCH("备选",N4))'], fill=PatternFill("solid", fgColor=AMBER)))

gray_headers = ["蛋白", "重复", "分组", "来源单元格", "原始灰度值", "主用原膜相对路径", "主文件状态", "定量追溯说明"]
gray_rows = [[x["target"], f"rep{x['rep']}", x["group"], x["cell"], x["value"], x["path"], x["file_status"], x["trace"]]
             for x in gray_records]
ws = add_table_sheet("灰度单元格映射", "63 个原始灰度单元格—主用原膜映射",
                     "数值直接转录自 10.14灰度值(1).xlsx；未重新定量。具体 ROI、背景扣除和曝光选择仍需原始 ImageJ／实验记录。",
                     gray_headers, gray_rows,
                     {"A": 3, "B": 21, "C": 11, "D": 21, "E": 18, "F": 16, "G": 34, "H": 30, "I": 50})
for row in range(4, 67):
    ws.cell(row, 6).number_format = "0"

# Evidence records: six archive checks plus two discriminating PPT template matches.
evidence_headers = ["证据ID", "证据类型", "来源条目", "对应本地文件", "包内时间", "指标", "来源值", "本地值", "结果", "补充说明"]
evidence_rows = []
for i, item in enumerate(rar_rows, 1):
    evidence_rows.append([f"RAR-{i:02d}", "044 来源包 CRC32", item["entry"], item["local"], item["archive_time"],
                          "CRC32 / 字节数", f"{item['archive_crc32']} / {item['bytes']}",
                          f"{item['local_crc32']} / {item['bytes']}", "一致", "同时与 044.pptx 打包于 10144(1).rar"])
for i, item in enumerate(ppt_evidence, 1):
    evidence_rows.append([f"PPT-{i:02d}", "PPT 裁带模板匹配", item["entry"], item["local"], None,
                          "Zero-mean NCC", None, round(item["score"], 6), "图像对应确认",
                          f"同名次优候选 {item['next_candidate']}，NCC={item['next_score']:.6f}；匹配尺度={item['scale']:.4f}"])
ws = add_table_sheet("PPT与RAR证据", "PPT 与 RAR 来源证据",
                     "CRC32 用于确认包内条目与本地文件一致；NCC 用于确认 PPT 裁带的图像来源。二者均不替代原始 ImageJ ROI 记录。",
                     evidence_headers, evidence_rows,
                     {"A": 3, "B": 13, "C": 23, "D": 35, "E": 32, "F": 21, "G": 20, "H": 25, "I": 25, "J": 20, "K": 58})

source_headers = ["项目", "相对路径／说明", "SHA-256／口径"]
source_rows = [
    ["灰度工作簿", "06_source_material/wetlab_wb_raw_20261030/10.14灰度值(1).xlsx", sha256(SOURCE)],
    ["044 PPT", "06_source_material/wetlab_wb_raw_20261030/044.pptx", sha256(RAW / "044.pptx")],
    ["044 来源包", "06_source_material/wetlab_wb_raw_20261030/10144(1).rar", sha256(RAW / "10144(1).rar")],
    ["其余条带 PPT", "06_source_material/wetlab_wb_raw_20261030/022.pptx", sha256(RAW / "022.pptx")],
    ["图像分析记录", "07_audit/wb_source_mapping/image_provenance_analysis.json", sha256(AUDIT / "image_provenance_analysis.json")],
    ["分组顺序", "rep1=C:E；rep2=G:I；rep3=K:M", "Control / HKSA+Baicalein / HKSA"],
    ["内参", "原始表 actin 行；稿件与本清单统一为 β-actin", "不改写为 GAPDH"],
    ["主用文件确认", "044 来源包内 6 个 p65／p-p65 TIFF，或其他重复／靶标的唯一全膜", "文件级对应；不代表 ROI 级复现"],
    ["备选版本", "未进入 044 来源包的同名 p65／p-p65 文件", "保留归档；不作现稿灰度来源指认"],
    ["定量追溯", "当前归档未发现 ImageJ ROI 集、测量输出或背景扣除记录", "需原始记录才能完成 ROI 级确证"],
    ["外部来源 URL", "无；全部使用本地归档原始实验文件", "不引入外部数据"],
]
ws = add_table_sheet("来源与口径", "来源文件与状态定义",
                     "本工作簿只解决 Source Data 文件对应与可追溯性，不重算 WB 比率、统计或图件。",
                     source_headers, source_rows, {"A": 3, "B": 25, "C": 92, "D": 76})

for sheet in wb.worksheets:
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUT)

print("INDEPENDENT_QA")
print(f"TIFF total={len(records)}")
print(f"Primary={sum(x['role'] == '主用原膜' for x in records)}")
print(f"p65/p-p65 primary={sum(x['role'] == '主用原膜' and x['target'] in {'p65','p-p65'} for x in records)}")
print(f"Alternates={sum(x['role'] == '备选版本' for x in records)}")
print(f"RAR CRC32 matches={sum(x['match'] == '一致' for x in rar_rows)}")
print(f"Grey cells={len(gray_records)}")
for item in ppt_evidence:
    print(f"{item['entry']} -> {item['local']}: NCC={item['score']:.6f}; next={item['next_score']:.6f}")
print(f"WORKBOOK={OUT}")
