# Public source snapshot: replace path placeholders before running.
from pathlib import Path
import json
import math
import zipfile

import openpyxl
from PIL import Image
from pypdf import PdfReader

BASE = Path(r"<PROJECT_ROOT>\outputs\elisa_redraw_20261026")
XLSX = BASE / "ELISA三因子_9组重算与制图_20261026.xlsx"
RESULTS = json.loads((BASE / "elisa_recalculated_20261026.json").read_text(encoding="utf-8"))

wbf = openpyxl.load_workbook(XLSX, data_only=False)
wbv = openpyxl.load_workbook(XLSX, data_only=True)
max_delta = {"OD": 0.0, "concentration": 0.0, "mean": 0.0, "SD": 0.0,
             "raw_P": 0.0, "adjusted_P": 0.0, "ANOVA_F": 0.0, "ANOVA_P": 0.0}
checks = 0


def compare(kind, actual, expected, rel_tol=1e-8):
    global checks
    if not isinstance(actual, (int, float)) or not math.isfinite(float(actual)):
        raise AssertionError(f"{kind}: invalid cached value {actual!r}; expected {expected}")
    delta = abs(float(actual) - float(expected))
    max_delta[kind] = max(max_delta[kind], delta)
    if not math.isclose(float(actual), float(expected), rel_tol=rel_tol, abs_tol=1e-300):
        raise AssertionError(f"{kind}: actual={actual!r}, expected={expected!r}, delta={delta}")
    checks += 1


for analyte, bundle in RESULTS.items():
    ws = wbv[analyte]
    analysis = bundle["analysis"]
    groups = list(analysis["groups"])
    for gi, group in enumerate(groups, start=2):
        expected_group = analysis["groups"][group]
        raw_group = bundle["plate"]["groups"][group]
        for rep in range(5):
            compare("OD", ws.cell(row=7 + rep, column=gi).value, raw_group["od"][rep])
            compare("concentration", ws.cell(row=14 + rep, column=gi).value,
                    expected_group["values"][rep])
        compare("mean", ws.cell(row=20, column=gi).value, expected_group["mean"])
        compare("SD", ws.cell(row=21, column=gi).value, expected_group["sd"])
        compare("concentration", ws.cell(row=22, column=gi).value, expected_group["n"])
    comparisons = list(analysis["comparisons"].values())
    for row, expected in zip(range(27, 32), comparisons):
        compare("raw_P", ws.cell(row=row, column=4).value, expected["p_raw"], rel_tol=1e-6)
        compare("adjusted_P", ws.cell(row=row, column=5).value,
                expected["p_bonferroni"], rel_tol=1e-6)
        actual_star = ws.cell(row=row, column=6).value
        if actual_star != expected["significance"]:
            raise AssertionError(f"{analyte} row {row}: stars {actual_star!r} != {expected['significance']!r}")
        checks += 1
    compare("ANOVA_F", ws["I32"].value, analysis["anova"]["F"], rel_tol=1e-7)
    compare("ANOVA_P", ws["I33"].value, analysis["anova"]["p"], rel_tol=1e-6)
    if not analysis["dose_strictly_decreasing"]:
        raise AssertionError(f"{analyte}: dose sequence is not strictly decreasing")
    occlusion = analysis["comparisons"]["HKSA+IKK-16 vs HKSA+BAI(24)+IKK-16"]
    if occlusion["significance"] != "ns":
        raise AssertionError(f"{analyte}: occlusion comparison is not ns")
    print(f"{analyte}: 9 groups × 5 replicates; dose 6→48 μM strictly decreasing; occlusion ns; "
          f"BAI24 inhibition={analysis['bai24_induction_normalized_inhibition_percent']:.3f}%")

formula_count = 0
missing_cache = []
error_values = []
for wsf in wbf.worksheets:
    wsv = wbv[wsf.title]
    for row in wsf.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formula_count += 1
                cached = wsv[cell.coordinate].value
                if cached is None:
                    missing_cache.append(f"{wsf.title}!{cell.coordinate}")
                if isinstance(cached, str) and cached.startswith("#"):
                    error_values.append(f"{wsf.title}!{cell.coordinate}={cached}")

chart_count = sum(len(ws._charts) for ws in wbf.worksheets)
image_count = sum(len(ws._images) for ws in wbf.worksheets)
if formula_count != 285 or missing_cache or error_values:
    raise AssertionError(f"formula_count={formula_count}, missing={missing_cache[:10]}, errors={error_values[:10]}")
if chart_count != 3 or image_count < 1:
    raise AssertionError(f"charts={chart_count}, images={image_count}")
with zipfile.ZipFile(XLSX) as archive:
    bad_member = archive.testzip()
    if bad_member:
        raise AssertionError(f"Corrupt XLSX member: {bad_member}")
png = Image.open(BASE / "Fig_ELISA_9groups_20261026.png")
pdf = PdfReader(str(BASE / "Fig_ELISA_9groups_20261026.pdf"))
if len(pdf.pages) != 1:
    raise AssertionError("PDF page count is not 1")
print(f"Cross-checks passed: {checks}")
print("Maximum absolute deltas:", max_delta)
print(f"Workbook: {formula_count} formulas, all cached; no formula errors; {chart_count} native charts; {image_count} embedded image")
print(f"PNG: {png.size}; PDF pages: {len(pdf.pages)}; XLSX ZIP integrity: OK")
